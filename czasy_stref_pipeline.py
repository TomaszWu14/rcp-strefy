"""
Pipeline: Obliczanie czasu pracy w strefach magazynowych
Wejście : Czasy.xlsx (eksport SAP EWM - zakładka Arkusz1)
Wyjście : czasy_stref_wynik.xlsx

Logika obliczania czasu (SAP EWM Transfer Order):
  Każde zadanie ma 2 pozycje:
    poz1: potwierdzenie WYJŚCIA ze strefy  → dt_exit
    poz2: potwierdzenie WEJŚCIA do strefy  → dt_enter
  Czas w strefie X = dt następnego poz1 tego samego HU - dt poz2 wejścia do X
  Obsługuje przypadki:
    - brak odbicia na koniec zmiany (odrzuca rekord bez następnego poz1)
    - przejście przez północ / brak odbicia kilka dni (limit MAX_CZAS_H)
    - pracownicy którzy nie odbijają STREFA_KONIEC
"""

import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INPUT_FILE  = Path(__file__).parent / "Czasy.xlsx"
OUTPUT_FILE = Path(__file__).parent / "czasy_stref_wynik.xlsx"

STREFY_ROBOCZE = [
    "STREFA_PICKING", "STREFA_PRZYJĘCIA", "STREFA_PRZEWOZNIK_A", "STREFA_PRZEWOZNIK_B",
    "STREFA_BUS", "STREFA_EXPORT", "STREFA_INV", "STREFA_ZWROTY",
    "STREFA_INNE", "STREFA_KONTENERY", "STREFA_REO", "STREFA_KLIENT_A",
    "STREFA_ADMIN_MAG", "STREFA_SZKOLENIE", "STREFA_ODDZIAŁY",
    "STREFA_ODDZIAŁ_KPL", "STREFA_ODDZIAŁ_PAK", "STREFA_KLIENT_B",
    "STREFA_ZERO", "STREFA_REKLAMACJE", "STREFA_KP_ELEKTRYK",
    "STREFA_KP_PALECIAK", "STREFA_PRZERWA", "STREFA_UKRAINA", "STREFA_KLIENT_C",
]
MAX_CZAS_H = 12   # odrzuć odcinki >12h (brak odbicia przez kilka zmian)

# ── 1. WCZYTAJ ────────────────────────────────────────────────────────────────
print("1/5  Wczytywanie danych...")
raw = pd.read_excel(INPUT_FILE, sheet_name="Arkusz1", dtype=str)

raw = raw[[c for c in raw.columns if "UniqueName" not in c]].copy()
raw.columns = [
    "data_utworzenia", "zadanie", "hu", "strefa_zrodlowa", "pozycja",
    "data_potwierdzenia", "czas_potwierdzenia", "status", "strefa_docelowa",
]

raw = raw[raw["status"] == "Potwierdzone"].copy()
raw["dt"] = pd.to_datetime(
    raw["data_potwierdzenia"] + " " + raw["czas_potwierdzenia"],
    format="%d.%m.%Y %H:%M:%S", errors="coerce")
raw = raw.dropna(subset=["dt"])
print(f"   Wierszy: {len(raw):,}")

# ── 2. OBLICZ CZAS W STREFACH ────────────────────────────────────────────────
print("2/5  Obliczanie czasu...")

p2 = (raw[raw["pozycja"] == "2"]
      .drop_duplicates(subset=["hu", "dt"])
      .sort_values(["hu", "dt"])
      .reset_index(drop=True))

p1 = (raw[raw["pozycja"] == "1"][["hu", "dt"]]
      .drop_duplicates(subset=["hu", "dt"])
      .rename(columns={"dt": "dt_exit"})
      .sort_values(["hu", "dt_exit"])
      .reset_index(drop=True))

# Per-group merge_asof: dla każdego wejścia (poz2) znajdź najbliższy następny
# poz1 tego samego HU → to jest moment wyjścia ze strefy
results = []
for hu, g2 in p2.groupby("hu"):
    g1 = p1[p1["hu"] == hu].reset_index(drop=True)
    if g1.empty:
        continue
    merged = pd.merge_asof(
        g2.reset_index(drop=True),
        g1[["dt_exit"]],
        left_on="dt", right_on="dt_exit",
        direction="forward")
    results.append(merged)

df = pd.concat(results, ignore_index=True)
df["czas_s"] = (df["dt_exit"] - df["dt"]).dt.total_seconds()
df = df[df["strefa_docelowa"].isin(STREFY_ROBOCZE)]
df = df[(df["czas_s"] > 0) & (df["czas_s"] <= MAX_CZAS_H * 3600)]
df["data"] = pd.to_datetime(df["data_potwierdzenia"], format="%d.%m.%Y", errors="coerce")
df["pracownik"] = df["hu"].str.replace("^HU_", "", regex=True).str.replace("_", " ")
df["rok_miesiac"] = df["data"].dt.to_period("M")
df["czas_h"] = (df["czas_s"] / 3600).round(4)

print(f"   Rekordów: {len(df):,} | Pracownicy: {df['pracownik'].nunique()}")
print(f"   Zakres: {df['data'].min().date()} – {df['data'].max().date()}")

# ── 3. PIVOTY ─────────────────────────────────────────────────────────────────
print("3/5  Tworzę pivoty...")

def s_to_hms(s):
    s = int(s)
    return f'{s//3600:02}:{(s%3600)//60:02}:{s%60:02}'

# ── Dzienny decimal ───────────────────────────────────────────────────────────
daily_h = df.groupby(["data", "pracownik", "strefa_docelowa"])["czas_h"].sum().reset_index()
pivot_d = (daily_h.pivot_table(index=["data", "pracownik"], columns="strefa_docelowa",
           values="czas_h", aggfunc="sum", fill_value=0).reset_index())
pivot_d.columns.name = None
sc = [c for c in pivot_d.columns if c.startswith("STREFA_")]
pivot_d["SUMA_h"] = pivot_d[sc].sum(axis=1).round(4)
pivot_d["data"] = pivot_d["data"].dt.date

# ── Dzienny GG:MM:SS ──────────────────────────────────────────────────────────
daily_s = df.groupby(["data", "pracownik", "strefa_docelowa"])["czas_s"].sum().reset_index()
pivot_ds = (daily_s.pivot_table(index=["data", "pracownik"], columns="strefa_docelowa",
            values="czas_s", aggfunc="sum", fill_value=0).reset_index())
pivot_ds.columns.name = None
sc2 = [c for c in pivot_ds.columns if c.startswith("STREFA_")]
pivot_ds["SUMA_s"] = pivot_ds[sc2].sum(axis=1)
pivot_dhms = pivot_ds[["data", "pracownik"]].copy()
pivot_dhms["data"] = pivot_ds["data"].dt.date
for col in sc2:
    pivot_dhms[col] = pivot_ds[col].apply(s_to_hms)
pivot_dhms["SUMA"] = pivot_ds["SUMA_s"].apply(s_to_hms)

# ── Miesięczny decimal ────────────────────────────────────────────────────────
monthly_h = df.groupby(["rok_miesiac", "pracownik", "strefa_docelowa"])["czas_h"].sum().reset_index()
pivot_m = (monthly_h.pivot_table(index=["rok_miesiac", "pracownik"], columns="strefa_docelowa",
           values="czas_h", aggfunc="sum", fill_value=0).reset_index())
pivot_m.columns.name = None
scm = [c for c in pivot_m.columns if c.startswith("STREFA_")]
pivot_m["SUMA_h"] = pivot_m[scm].sum(axis=1).round(4)
pivot_m["rok_miesiac"] = pivot_m["rok_miesiac"].astype(str)

# ── Miesięczny GG:MM:SS ───────────────────────────────────────────────────────
monthly_s = df.groupby(["rok_miesiac", "pracownik", "strefa_docelowa"])["czas_s"].sum().reset_index()
pivot_ms = (monthly_s.pivot_table(index=["rok_miesiac", "pracownik"], columns="strefa_docelowa",
            values="czas_s", aggfunc="sum", fill_value=0).reset_index())
pivot_ms.columns.name = None
scms = [c for c in pivot_ms.columns if c.startswith("STREFA_")]
pivot_ms["SUMA_s"] = pivot_ms[scms].sum(axis=1)
pivot_mhms = pivot_ms[["rok_miesiac", "pracownik"]].copy()
pivot_mhms["rok_miesiac"] = pivot_ms["rok_miesiac"].astype(str)
for col in scms:
    pivot_mhms[col] = pivot_ms[col].apply(s_to_hms)
pivot_mhms["SUMA"] = pivot_ms["SUMA_s"].apply(s_to_hms)

# ── Dane surowe ───────────────────────────────────────────────────────────────
surowe = df[["data", "pracownik", "hu", "strefa_docelowa", "dt", "dt_exit", "czas_s"]].copy()
surowe["czas"] = pd.to_timedelta(surowe["czas_s"], unit="s").apply(
    lambda x: f'{int(x.total_seconds()//3600):02}:{int((x.total_seconds()%3600)//60):02}:{int(x.total_seconds()%60):02}')
surowe = surowe.rename(columns={"dt": "wejscie", "dt_exit": "wyjscie", "strefa_docelowa": "strefa"})
surowe["data"] = surowe["data"].dt.date
surowe = surowe.drop(columns=["czas_s"])

# ── 4. ZAPIS ──────────────────────────────────────────────────────────────────
print("4/5  Zapisuję plik...")
with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as w:
    pivot_d.to_excel(w,    sheet_name="Dzienny_h",         index=False)
    pivot_dhms.to_excel(w, sheet_name="Dzienny_ggmmss",    index=False)
    pivot_m.to_excel(w,    sheet_name="Miesieczny_h",      index=False)
    pivot_mhms.to_excel(w, sheet_name="Miesieczny_ggmmss", index=False)
    surowe.to_excel(w,     sheet_name="Dane_surowe",        index=False)

print("5/5  Formatuję...")
import time; time.sleep(0.5)
wb = load_workbook(OUTPUT_FILE)
H_FILL = PatternFill("solid", fgColor="1F4E79")
S_FILL = PatternFill("solid", fgColor="2E75B6")
G_FILL = PatternFill("solid", fgColor="375623")
A_FILL = PatternFill("solid", fgColor="DEEAF1")
W_FILL = PatternFill("solid", fgColor="FFFFFF")
HF  = Font(bold=True, color="FFFFFF", name="Arial", size=9)
BF  = Font(name="Arial", size=9)
BLD = Font(bold=True, name="Arial", size=9)
CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
LFT = Alignment(horizontal="left",   vertical="center")
T   = Side(style="thin", color="BFBFBF")
BDR = Border(left=T, right=T, top=T, bottom=T)

def fmt(ws, fc=2, zero_val=None):
    for c in ws[1]:
        c.fill = (S_FILL if (c.value and str(c.value).startswith("STREFA"))
                  else (G_FILL if c.value in ("SUMA_h", "SUMA") else H_FILL))
        c.font = HF; c.alignment = CTR; c.border = BDR
    for ri, row in enumerate(ws.iter_rows(min_row=2), 2):
        alt = ri % 2 == 0
        for c in row:
            c.border = BDR
            if c.column <= fc:
                c.font = BLD; c.alignment = LFT
            else:
                c.font = BF; c.alignment = CTR
                c.fill = A_FILL if alt else W_FILL
                if zero_val is not None and c.value == zero_val:
                    c.value = None
    for col in ws.columns:
        cl = get_column_letter(col[0].column)
        ml = max((len(str(c.value or "")) for c in col), default=6)
        ws.column_dimensions[cl].width = (min(ml + 2, 22) if col[0].column <= fc
                                           else max(min(ml + 1, 12), 8))
    ws.freeze_panes = ws.cell(row=2, column=fc + 1)
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 36

fmt(wb["Dzienny_h"],         zero_val=0.0)
fmt(wb["Dzienny_ggmmss"],    zero_val="00:00:00")
fmt(wb["Miesieczny_h"],      zero_val=0.0)
fmt(wb["Miesieczny_ggmmss"], zero_val="00:00:00")

wb["Dzienny_h"].sheet_properties.tabColor         = "2E75B6"
wb["Dzienny_ggmmss"].sheet_properties.tabColor    = "70AD47"
wb["Miesieczny_h"].sheet_properties.tabColor      = "2E75B6"
wb["Miesieczny_ggmmss"].sheet_properties.tabColor = "70AD47"

ws_s = wb["Dane_surowe"]
for c in ws_s[1]:
    c.fill = H_FILL; c.font = HF; c.alignment = CTR; c.border = BDR
for row in ws_s.iter_rows(min_row=2):
    for c in row:
        c.font = BF; c.alignment = LFT; c.border = BDR
for col in ws_s.columns:
    ws_s.column_dimensions[get_column_letter(col[0].column)].width = min(
        max((len(str(c.value or "")) for c in col), default=6) + 2, 28)
ws_s.freeze_panes = "A2"
ws_s.auto_filter.ref = ws_s.dimensions

wb.save(OUTPUT_FILE)
print(f"\n✓ Gotowe → {OUTPUT_FILE}")
print(f"  Dzienny_h / Dzienny_ggmmss: {len(pivot_d):,} wierszy")
print(f"  Miesieczny_h / Miesieczny_ggmmss: {len(pivot_m):,} wierszy")
print(f"  Dane_surowe: {len(surowe):,} wierszy")
